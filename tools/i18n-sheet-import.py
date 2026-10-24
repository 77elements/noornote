#!/usr/bin/env python3
"""Import a filled XLSX translation sheet into a .po catalog (WP13, D6).

Counterpart to i18n-sheet-export.py — see docs/todos/internationalization.md
§7.7 for the runbook. Validates every filled row before touching the catalog:

  - placeholder check: `%s` / `%N$s` count must match the English source
  - plural rows need BOTH forms filled
  - formula cells (`=…`) are rejected — read as text, never evaluated
  - unknown keys (po changed since export) are skipped with a warning
  - duplicates within the sheet are rejected

Merge rules (gettext-safe):
  - empty msgstr        → filled (marked fuzzy only with --fuzzy)
  - fuzzy msgstr        → overwritten AND un-fuzzied (MT draft confirmed/fixed)
  - final msgstr        → NEVER overwritten (counted as `already-final`)

Usage:
    python3 tools/i18n-sheet-import.py --sheet noornote-translate-ar.xlsx \
        --lang ar [--fuzzy] [--compile]

Exit codes: 0 = something imported, 1 = nothing imported / fatal.
"""

from __future__ import annotations

import argparse
import re
import subprocess
import sys
from pathlib import Path

import polib  # pip3 install --user --break-system-packages polib
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
PO_DIR = ROOT / "src" / "i18n" / "po"

PLACEHOLDER_RE = re.compile(r"%(?:\d+\$)?s")

# sheet columns (see export docstring)
COL_KEY, COL_CTX, COL_FORM = "A", "B", "C"
COL_EN, COL_EN_PL = "D", "E"
COL_FILL, COL_FILL_PL = "J", "K"


def placeholder_signature(text: str) -> tuple[int, frozenset[str]]:
    """(number of sequential %s, set of explicit %N$s indices)."""
    matches = PLACEHOLDER_RE.findall(text)
    explicit = {m for m in matches if m != "%s"}
    sequential = matches.count("%s")
    return sequential, frozenset(explicit)


def placeholders_match(source: str, translation: str) -> bool:
    return placeholder_signature(source) == placeholder_signature(translation)


def read_rows(sheet_path: Path) -> list[dict[str, str]]:
    """Read data rows (row 2+) as plain text. data_only=False keeps formulas
    visible as `=…` strings so they can be rejected instead of evaluated."""
    wb = load_workbook(sheet_path, data_only=False, read_only=True)
    ws = wb.active
    rows: list[dict[str, str]] = []
    for row in ws.iter_rows(min_row=2):
        cells = list(row) + [None] * (11 - len(row))  # read_only trims trailing cells
        get = lambda col: (str(cells[ord(col) - 65].value) if cells[ord(col) - 65].value is not None else "")
        record = {col: get(col).strip() for col in "ABCDEFGHIJK"}
        if record[COL_KEY]:
            rows.append(record)
    wb.close()
    return rows


def fill_entry(
    entry: polib.POEntry,
    sg: str,
    pl: str,
    is_plural_row: bool,
    fuzzy: bool,
) -> str:
    """Apply one validated row. Returns 'filled' | 'unfuzzied' (MT draft
    that got confirmed/fixed — only counts as unfuzzied when it actually
    lost its fuzzy flag)."""
    was_fuzzy = "fuzzy" in entry.flags
    if was_fuzzy:
        entry.flags.remove("fuzzy")
    if fuzzy:
        entry.flags.append("fuzzy")
    if is_plural_row:
        entry.msgstr_plural = {0: sg, 1: pl}
    else:
        entry.msgstr = sg
    return "unfuzzied" if was_fuzzy and not fuzzy else "filled"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--sheet", required=True, help="filled XLSX from the translator")
    ap.add_argument("--lang", required=True, help="target language code (e.g. ar)")
    ap.add_argument("--fuzzy", action="store_true", help="mark imported entries as fuzzy (MT drafts)")
    ap.add_argument("--compile", action="store_true", help="run i18n:compile after a successful import")
    args = ap.parse_args()

    if args.lang == "en":
        print("en is the source catalog — never import into en")
        return 1

    po_path = PO_DIR / f"{args.lang}.po"
    if not po_path.exists():
        print(f"no catalog at src/i18n/po/{args.lang}.po — run the export first (it bootstraps)")
        return 1
    po = polib.pofile(str(po_path), wrap_width=79)

    rows = read_rows(Path(args.sheet))
    if not rows:
        print("sheet has no data rows")
        return 1

    filled: list[str] = []
    unfuzzied: list[str] = []
    skipped_final: list[str] = []
    skipped_empty = 0
    invalid: list[str] = []
    seen: set[tuple[str | None, str]] = set()

    for no, row in enumerate(rows, start=2):  # XLSX row number for the report
        key, ctx, form = row[COL_KEY], row[COL_CTX] or None, row[COL_FORM]
        label = f"{ctx}:{key}" if ctx else key
        sheet_row = f"row {no} ({label})"

        if (ctx, key) in seen:
            invalid.append(f"{sheet_row}: duplicate key in sheet")
            continue
        seen.add((ctx, key))

        if row[COL_EN] != key:
            invalid.append(f"{sheet_row}: English source edited (must stay identical)")
            continue

        sg = row[COL_FILL]
        pl = row[COL_FILL_PL]
        if sg.startswith("=") or pl.startswith("="):
            invalid.append(f"{sheet_row}: formula cell — enter plain text, not a formula")
            continue
        if not sg and not pl:
            skipped_empty += 1
            continue

        entry = po.find(key, msgctxt=ctx)
        if entry is None or entry.obsolete:
            invalid.append(f"{sheet_row}: unknown key (catalog changed since export?)")
            continue

        is_plural_entry = bool(entry.msgid_plural)
        if is_plural_entry:
            if form != "Plural":
                invalid.append(f"{sheet_row}: marked Singular but the catalog expects a plural entry")
                continue
            if not sg or not pl:
                invalid.append(f"{sheet_row}: plural entry needs BOTH forms filled")
                continue
            if not placeholders_match(row[COL_EN], sg):
                invalid.append(f"{sheet_row}: singular %s placeholders don't match English")
                continue
            if not placeholders_match(row[COL_EN_PL], pl):
                invalid.append(f"{sheet_row}: plural %s placeholders don't match English plural")
                continue
        else:
            if form != "Singular":
                invalid.append(f"{sheet_row}: marked Plural but the catalog expects a singular entry")
                continue
            if pl:
                invalid.append(f"{sheet_row}: plural cell filled on a singular entry")
                continue
            if not placeholders_match(row[COL_EN], sg):
                invalid.append(f"{sheet_row}: %s placeholders don't match English")
                continue

        forms = entry.msgstr_plural or {}
        is_final = (
            not entry.msgid_plural and entry.msgstr.strip()
        ) or (
            entry.msgid_plural
            and forms.get(0, "").strip()
            and forms.get(1, "").strip()
        )
        if is_final and "fuzzy" not in entry.flags:
            skipped_final.append(label)
            continue
        outcome = fill_entry(entry, sg, pl, is_plural_entry, args.fuzzy)
        (unfuzzied if outcome == "unfuzzied" else filled).append(label)

    print(f"filled:          {len(filled)}")
    print(f"unfuzzied (MT):  {len(unfuzzied)}")
    print(f"left empty:      {skipped_empty}")
    print(f"already final:   {len(skipped_final)} (never overwritten)")
    if invalid:
        print(f"INVALID ({len(invalid)}) — fix these and re-import:")
        for line in invalid:
            print(f"  ✗ {line}")

    if filled or unfuzzied:
        po.save(str(po_path))
        print(f"saved: src/i18n/po/{args.lang}.po")
        if args.compile:
            subprocess.run(["bun", "run", "i18n:compile"], check=True)
        return 0

    print("nothing imported")
    return 1


if __name__ == "__main__":
    sys.exit(main())
